import pandas as pd
from datetime import datetime, timedelta
import os
from typing import Dict, List, Optional
import logging
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.utils.dataframe import dataframe_to_rows
import xlsxwriter

import sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '../..'))

from app.db.adaptive_dao import adaptive_dao

logger = logging.getLogger(__name__)


class ExcelReportGenerator:
    def __init__(self):
        self.dao = adaptive_dao
        
        # Colores de la academia (profesional)
        self.colors = {
            'header': 'FF1E3A8A',  # Azul oscuro profesional
            'subheader': 'FF3B82F6',  # Azul medio
            'accent': 'FF60A5FA',  # Azul claro
            'success': 'FF10B981',  # Verde
            'warning': 'FFF59E0B',  # Naranja
            'danger': 'FFEF4444',  # Rojo
            'light': 'FFF3F4F6'  # Gris claro
        }
    
    def generate_full_report(self, output_path: str = None, period_days: int = 30) -> str:
        """
        Genera reporte Excel completo con múltiples hojas y gráficos.
        """
        try:
            logger.info(f"📊 Generando reporte Excel para últimos {period_days} días...")
            
            # Definir archivo de salida
            if not output_path:
                timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
                output_path = f"reporte_jeanacademy_{timestamp}.xlsx"
            
            # Obtener datos
            data = self._fetch_report_data(period_days)
            
            # Crear Excel con xlsxwriter para más control
            with pd.ExcelWriter(output_path, engine='xlsxwriter') as writer:
                workbook = writer.book
                
                # Formatos globales
                formats = self._create_formats(workbook)
                
                # 1. Hoja de Resumen Ejecutivo
                self._create_summary_sheet(writer, data, formats)
                
                # 2. Hoja de Módulos
                self._create_modules_sheet(writer, data, formats)
                
                # 3. Hoja de Estudiantes
                self._create_students_sheet(writer, data, formats)
                
                # 4. Hoja de Entregas Detalladas
                self._create_submissions_sheet(writer, data, formats)
                
                # 5. Hoja de Estadísticas
                self._create_statistics_sheet(writer, data, formats)
                
                # Ajustar tamaños de columnas
                self._adjust_column_widths(writer)
            
            logger.info(f"✅ Reporte generado: {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"❌ Error generando reporte Excel: {e}")
            raise
    
    def _fetch_report_data(self, period_days: int) -> Dict:
        """Obtiene todos los datos necesarios para el reporte."""
        logger.info("📥 Obteniendo datos de la base de datos...")
        
        data = {}
        
        with self.dao.get_cursor() as cur:
            # Fecha de corte
            cutoff_date = datetime.now() - timedelta(days=period_days)
            
            # 1. Resumen general
            cur.execute("""
                SELECT 
                    COUNT(DISTINCT m.id) as total_modules,
                    COUNT(DISTINCT s.id) as total_students,
                    COUNT(DISTINCT sub.id) as total_submissions,
                    COUNT(DISTINCT CASE 
                        WHEN sub.detected_at > %s THEN sub.student_id 
                    END) as active_students
                FROM modules m
                LEFT JOIN submissions sub ON m.id = sub.module_id
                LEFT JOIN students s ON sub.student_id = s.id
            """, (cutoff_date,))
            data['summary'] = cur.fetchone()
            
            # 2. Datos por módulo
            cur.execute("""
                SELECT 
                    m.code,
                    m.name as module_name,
                    COUNT(DISTINCT sub.student_id) as students_count,
                    COUNT(sub.id) as submissions_count,
                    MAX(sub.detected_at AT TIME ZONE 'UTC')::timestamp as last_submission,
                    m.drive_folder_url
                FROM modules m
                LEFT JOIN submissions sub ON m.id = sub.module_id
                GROUP BY m.id, m.code, m.name, m.drive_folder_url
                ORDER BY m.order_index
            """)
            data['modules'] = cur.fetchall()
            
            # 3. Datos por estudiante
            cur.execute("""
                SELECT 
                    s.full_name,
                    s.email,
                    COUNT(DISTINCT sub.module_id) as modules_completed,
                    COUNT(sub.id) as total_submissions,
                    MAX(sub.detected_at AT TIME ZONE 'UTC')::timestamp as last_activity,
                    MIN(sub.detected_at AT TIME ZONE 'UTC')::timestamp as first_activity
                FROM students s
                LEFT JOIN submissions sub ON s.id = sub.student_id
                GROUP BY s.id, s.full_name, s.email
                ORDER BY modules_completed DESC, s.full_name
            """)
            data['students'] = cur.fetchall()
            
            # 4. Entregas recientes (con email para mejor contexto)
            cur.execute("""
                SELECT 
                    s.full_name as student_name,
                    s.email,
                    m.name as module_name,
                    sub.filename,
                    sub.file_extension,
                    sub.size_mb,
                    (sub.detected_at AT TIME ZONE 'UTC')::timestamp as detected_at,
                    sub.drive_url,
                    CASE 
                        WHEN s.email LIKE '%%@gmail.com' THEN '✅ Gmail'
                        WHEN s.email LIKE '%%.temp' THEN '⚠️ Temporal'
                        WHEN s.email LIKE '%%@%%' THEN '✅ Email real'
                        ELSE '❓ Revisar'
                    END as email_status
                FROM submissions sub
                JOIN students s ON sub.student_id = s.id
                JOIN modules m ON sub.module_id = m.id
                WHERE sub.detected_at > %s
                ORDER BY sub.detected_at DESC
                LIMIT 100
            """, (cutoff_date,))
            data['recent_submissions'] = cur.fetchall()
            
            # 5. Estadísticas por día
            cur.execute("""
                SELECT 
                    DATE(detected_at AT TIME ZONE 'UTC') as date,
                    COUNT(*) as submissions,
                    COUNT(DISTINCT student_id) as active_students
                FROM submissions
                WHERE detected_at > %s
                GROUP BY DATE(detected_at AT TIME ZONE 'UTC')
                ORDER BY date DESC
            """, (cutoff_date,))
            data['daily_stats'] = cur.fetchall()
            
            # 6. Top estudiantes más activos
            cur.execute("""
                SELECT 
                    s.full_name,
                    COUNT(sub.id) as submissions
                FROM students s
                JOIN submissions sub ON s.id = sub.student_id
                WHERE sub.detected_at > %s
                GROUP BY s.id, s.full_name
                ORDER BY submissions DESC
                LIMIT 10
            """, (cutoff_date,))
            data['top_students'] = cur.fetchall()
            
        
        logger.info(f"✅ Datos obtenidos: {len(data)} conjuntos")
        return data
    
    def _create_formats(self, workbook) -> Dict:
        """Crea formatos reutilizables para el Excel."""
        return {
            'title': workbook.add_format({
                'font_size': 20,
                'bold': True,
                'font_color': self.colors['header'],
                'align': 'center',
                'valign': 'vcenter'
            }),
            'subtitle': workbook.add_format({
                'font_size': 14,
                'bold': True,
                'font_color': self.colors['subheader'],
                'align': 'left'
            }),
            'header': workbook.add_format({
                'bold': True,
                'bg_color': self.colors['header'],
                'font_color': 'white',
                'border': 1,
                'align': 'center',
                'valign': 'vcenter',
                'text_wrap': True
            }),
            'cell': workbook.add_format({
                'border': 1,
                'align': 'left',
                'valign': 'vcenter'
            }),
            'number': workbook.add_format({
                'border': 1,
                'align': 'right',
                'num_format': '#,##0'
            }),
            'percentage': workbook.add_format({
                'border': 1,
                'align': 'right',
                'num_format': '0.0%'
            }),
            'date': workbook.add_format({
                'border': 1,
                'align': 'center',
                'num_format': 'dd/mm/yyyy'
            }),
            'datetime': workbook.add_format({
                'border': 1,
                'align': 'center',
                'num_format': 'dd/mm/yyyy hh:mm'
            }),
            'success': workbook.add_format({
                'bg_color': self.colors['success'],
                'font_color': 'white',
                'bold': True,
                'border': 1,
                'align': 'center'
            }),
            'warning': workbook.add_format({
                'bg_color': self.colors['warning'],
                'font_color': 'white',
                'bold': True,
                'border': 1,
                'align': 'center'
            }),
            'danger': workbook.add_format({
                'bg_color': self.colors['danger'],
                'font_color': 'white',
                'bold': True,
                'border': 1,
                'align': 'center'
            })
        }
    
    def _create_summary_sheet(self, writer, data: Dict, formats: Dict):
        """Crea hoja de resumen ejecutivo con KPIs."""
        df_summary = pd.DataFrame([data['summary']])
        df_summary.to_excel(writer, sheet_name='Resumen', index=False, startrow=5)
        
        worksheet = writer.sheets['Resumen']
        workbook = writer.book
        
        # Título
        worksheet.merge_range('A1:F2', '🎓 REPORTE ACADÉMICO - JEAN ACADEMY', formats['title'])
        worksheet.merge_range('A3:F3', f"Fecha de generación: {datetime.now().strftime('%d/%m/%Y %H:%M')}", formats['subtitle'])
        
        # KPIs principales
        row = 7
        col = 0
        
        # Cards de métricas
        metrics = [
            ('Total Módulos', data['summary']['total_modules'], self.colors['header']),
            ('Total Estudiantes', data['summary']['total_students'], self.colors['subheader']),
            ('Total Entregas', data['summary']['total_submissions'], self.colors['success']),
            ('Estudiantes Activos', data['summary']['active_students'], self.colors['accent'])
        ]
        
        for i, (label, value, color) in enumerate(metrics):
            card_format = workbook.add_format({
                'bg_color': color,
                'font_color': 'white',
                'bold': True,
                'border': 2,
                'align': 'center',
                'valign': 'vcenter',
                'font_size': 12
            })
            
            worksheet.merge_range(row, col + i*2, row+1, col + i*2 + 1, '', card_format)
            worksheet.write(row, col + i*2, label, card_format)
            worksheet.write(row+1, col + i*2, value or 0, card_format)
        
        # Gráfico de estudiantes más activos
        if data['top_students']:
            row = 12
            worksheet.write(row, 0, 'Top 10 Estudiantes Más Activos', formats['subtitle'])
            
            row += 2
            for i, student in enumerate(data['top_students']):
                worksheet.write(row + i, 0, student['full_name'], formats['cell'])
                worksheet.write(row + i, 1, student['submissions'], formats['number'])
            
            # Crear gráfico de barras
            chart = workbook.add_chart({'type': 'bar'})
            chart.add_series({
                'name': 'Entregas',
                'categories': ['Resumen', row, 0, row + len(data['top_students']) - 1, 0],
                'values': ['Resumen', row, 1, row + len(data['top_students']) - 1, 1],
                'fill': {'color': self.colors['accent']}
            })
            chart.set_title({'name': 'Estudiantes Más Activos'})
            chart.set_x_axis({'name': 'Estudiante'})
            chart.set_y_axis({'name': 'Número de Entregas'})
            worksheet.insert_chart('D14', chart)
    
    def _create_modules_sheet(self, writer, data: Dict, formats: Dict):
        """Crea hoja con detalle de módulos."""
        if data['modules']:
            df_modules = pd.DataFrame(data['modules'])
            
            # Renombrar columnas para mejor presentación
            df_modules.columns = ['Código', 'Módulo', 'Estudiantes', 'Entregas', 'Última Actividad', 'URL Drive']
            
            # Escribir datos
            df_modules.to_excel(writer, sheet_name='Módulos', index=False, startrow=3)
            
            worksheet = writer.sheets['Módulos']
            
            # Título
            worksheet.merge_range('A1:F2', '📚 DETALLE DE MÓDULOS', formats['title'])
            
            # Formato de encabezados
            for col_num, col_name in enumerate(df_modules.columns):
                worksheet.write(3, col_num, col_name, formats['header'])
            
            # Formato de datos
            for row_num in range(len(df_modules)):
                for col_num in range(len(df_modules.columns)):
                    value = df_modules.iloc[row_num, col_num]
                    
                    if col_num in [2, 3]:  # Números
                        worksheet.write(row_num + 4, col_num, value or 0, formats['number'])
                    elif col_num == 4:  # Fecha
                        if value and pd.notna(value):
                            worksheet.write_datetime(row_num + 4, col_num, value, formats['datetime'])
                        else:
                            worksheet.write(row_num + 4, col_num, 'Sin actividad', formats['cell'])
                    else:
                        worksheet.write(row_num + 4, col_num, value or '', formats['cell'])
    
    def _create_students_sheet(self, writer, data: Dict, formats: Dict):
        """Crea hoja con detalle de estudiantes."""
        if data['students']:
            df_students = pd.DataFrame(data['students'])
            
            # Renombrar columnas
            df_students.columns = ['Nombre', 'Email', 'Módulos Completados', 'Total Entregas', 
                                  'Última Actividad', 'Primera Actividad']
            
            # Calcular progreso
            total_modules = data['summary']['total_modules'] or 1
            df_students['Progreso %'] = (df_students['Módulos Completados'] / total_modules * 100).round(1)
            
            # Escribir datos
            df_students.to_excel(writer, sheet_name='Estudiantes', index=False, startrow=3)
            
            worksheet = writer.sheets['Estudiantes']
            
            # Título
            worksheet.merge_range('A1:G2', '👥 DETALLE DE ESTUDIANTES', formats['title'])
            
            # Formato de encabezados
            for col_num, col_name in enumerate(df_students.columns):
                worksheet.write(3, col_num, col_name, formats['header'])
            
            # Formato de datos con colores según progreso
            for row_num in range(len(df_students)):
                progress = df_students.iloc[row_num, 6]  # Columna de Progreso %
                
                # Determinar formato según progreso
                if progress >= 80:
                    progress_format = formats['success']
                elif progress >= 50:
                    progress_format = formats['warning']
                else:
                    progress_format = formats['danger']
                
                for col_num in range(len(df_students.columns)):
                    value = df_students.iloc[row_num, col_num]
                    
                    if col_num in [2, 3]:  # Números
                        worksheet.write(row_num + 4, col_num, value or 0, formats['number'])
                    elif col_num in [4, 5]:  # Fechas
                        if value and pd.notna(value):
                            worksheet.write_datetime(row_num + 4, col_num, value, formats['datetime'])
                        else:
                            worksheet.write(row_num + 4, col_num, 'N/A', formats['cell'])
                    elif col_num == 6:  # Progreso
                        worksheet.write(row_num + 4, col_num, f"{value}%", progress_format)
                    else:
                        worksheet.write(row_num + 4, col_num, value or '', formats['cell'])
    
    def _create_submissions_sheet(self, writer, data: Dict, formats: Dict):
        """Crea hoja con entregas recientes."""
        if data['recent_submissions']:
            df_submissions = pd.DataFrame(data['recent_submissions'])
            
            # Renombrar columnas (ahora incluye email y estado del email)
            df_submissions.columns = ['Estudiante', 'Email', 'Módulo', 'Archivo', 'Tipo', 
                                     'Tamaño (MB)', 'Fecha Detección', 'URL Drive', 'Estado Email']
            
            df_submissions.to_excel(writer, sheet_name='Entregas Recientes', index=False, startrow=3)
            
            worksheet = writer.sheets['Entregas Recientes']
            
            # Título (ahora hasta columna I por las 9 columnas)
            worksheet.merge_range('A1:I2', '📝 ENTREGAS RECIENTES', formats['title'])
            
            # Formato de encabezados
            for col_num, col_name in enumerate(df_submissions.columns):
                worksheet.write(3, col_num, col_name, formats['header'])
    
    def _create_statistics_sheet(self, writer, data: Dict, formats: Dict):
        """Crea hoja con estadísticas y gráficos."""
        worksheet = writer.book.add_worksheet('Estadísticas')
        
        # Título
        worksheet.merge_range('A1:H2', '📈 ESTADÍSTICAS Y TENDENCIAS', formats['title'])
        
        # Actividad diaria
        if data['daily_stats']:
            row = 4
            worksheet.write(row, 0, 'Actividad Diaria', formats['subtitle'])
            
            row += 2
            worksheet.write(row, 0, 'Fecha', formats['header'])
            worksheet.write(row, 1, 'Entregas', formats['header'])
            worksheet.write(row, 2, 'Estudiantes Activos', formats['header'])
            
            for i, day_stat in enumerate(data['daily_stats']):
                worksheet.write(row + i + 1, 0, day_stat['date'], formats['date'])
                worksheet.write(row + i + 1, 1, day_stat['submissions'], formats['number'])
                worksheet.write(row + i + 1, 2, day_stat['active_students'], formats['number'])
            
            # Gráfico de líneas para actividad
            chart = writer.book.add_chart({'type': 'line'})
            chart.add_series({
                'name': 'Entregas',
                'categories': ['Estadísticas', row + 1, 0, row + len(data['daily_stats']), 0],
                'values': ['Estadísticas', row + 1, 1, row + len(data['daily_stats']), 1],
                'line': {'color': self.colors['header']}
            })
            chart.add_series({
                'name': 'Estudiantes Activos',
                'categories': ['Estadísticas', row + 1, 0, row + len(data['daily_stats']), 0],
                'values': ['Estadísticas', row + 1, 2, row + len(data['daily_stats']), 2],
                'line': {'color': self.colors['success']}
            })
            chart.set_title({'name': 'Tendencia de Actividad'})
            chart.set_x_axis({'name': 'Fecha'})
            chart.set_y_axis({'name': 'Cantidad'})
            worksheet.insert_chart('E6', chart, {'x_scale': 1.5, 'y_scale': 1.5})
    
    def _adjust_column_widths(self, writer):
        """Ajusta anchos de columna para mejor visualización."""
        for sheet_name in writer.sheets:
            worksheet = writer.sheets[sheet_name]
            
            # Anchos estándar por tipo de columna
            column_widths = {
                'A': 20,  # Nombres
                'B': 30,  # Emails/Descripciones
                'C': 15,  # Números
                'D': 15,  # Números
                'E': 20,  # Fechas
                'F': 20,  # URLs
                'G': 15,  # Porcentajes
                'H': 15   # Adicional
            }
            
            for col, width in column_widths.items():
                worksheet.set_column(f'{col}:{col}', width)


def test_excel_report():
    """Función de prueba para generar reporte Excel."""
    try:
        logger.info("🧪 Probando generación de reporte Excel...")
        
        generator = ExcelReportGenerator()
        output_file = generator.generate_full_report(period_days=30)
        
        logger.info(f"✅ Reporte generado exitosamente: {output_file}")
        
        # Verificar que el archivo existe
        if os.path.exists(output_file):
            file_size = os.path.getsize(output_file) / 1024  # KB
            logger.info(f"📊 Tamaño del archivo: {file_size:.2f} KB")
            return output_file
        else:
            logger.error("❌ El archivo no se creó")
            return None
            
    except Exception as e:
        logger.error(f"❌ Error en prueba: {e}")
        import traceback
        traceback.print_exc()
        return None


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    test_excel_report()